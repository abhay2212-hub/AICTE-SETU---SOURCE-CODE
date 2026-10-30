# Owner: Dev B
"""Reports router: role scoping, fixture/real switch, missing integrity score and file formats.

Real mode runs against a fake connection that answers each query by its SQL and records
the parameters, so scoping is checked without Postgres.
"""

from __future__ import annotations

import io
import json
import uuid
from contextlib import contextmanager
from pathlib import Path

import openpyxl
import pytest

from app.routers import reports as router
from tests.conftest import as_role

GROUPS = ("state", "programme", "category", "gender")
SHADOW = "10000000-0000-4000-8000-000000000003"
NARMADA = "10000000-0000-4000-8000-000000000002"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
OPENAPI = Path(__file__).resolve().parents[2] / "shared" / "openapi.json"

CSE, CE = uuid.uuid4(), uuid.uuid4()
ENROLMENT_ROWS = [
    {
        "programme_code": "UG-CSE",
        "state": "Uttar Pradesh",
        "category": "general",
        "gender": "male",
        "opted_out": False,
    },
    {
        "programme_code": "UG-CSE",
        "state": "Madhya Pradesh",
        "category": "obc",
        "gender": "female",
        "opted_out": True,
    },
    {
        "programme_code": "UG-CE",
        "state": "Uttar Pradesh",
        "category": "sc",
        "gender": "female",
        "opted_out": False,
    },
]
# Shadow UG-CSE: 62 regular + 3 per other pool = 71 active admissions against 60 seats.
SHADOW_ADMISSIONS = [
    {"programme_id": CSE, "programme_code": "UG-CSE", "category": "general", "has_apaar": i > 1}
    for i in range(71)
]
SHADOW_INTAKE = [
    {
        "programme_id": CSE,
        "programme_code": "UG-CSE",
        "pool": "regular",
        "approved_seats": 40,
        "rule_version": "APH-2025",
    },
    {
        "programme_id": CSE,
        "programme_code": "UG-CSE",
        "pool": "regular",
        "approved_seats": 51,
        "rule_version": "APH-2026",
    },
    *[
        {
            "programme_id": CSE,
            "programme_code": "UG-CSE",
            "pool": pool,
            "approved_seats": 3,
            "rule_version": "APH-2026",
        }
        for pool in ("tfw", "fn_oci_pio", "nri")
    ],
    # Seats but no admissions: still listed, admitted 0.
    {
        "programme_id": CE,
        "programme_code": "UG-CE",
        "pool": "regular",
        "approved_seats": 60,
        "rule_version": "APH-2026",
    },
]


class FakeCursor:
    def __init__(self, rows):
        self.rows = rows

    def fetchone(self):
        return self.rows[0] if self.rows else None

    def fetchall(self):
        return self.rows


class FakeConn:
    def __init__(self, score: float | None = 38.0):
        self.score = score
        self.calls: list[tuple[str, tuple]] = []

    def execute(self, sql: str, params: tuple = ()):
        self.calls.append((sql, params))
        if "from public.institutions where id" in sql:
            return FakeCursor(
                [
                    {
                        "id": uuid.UUID(params[0]),
                        "aicte_code": "1-9999003",
                        "name": "Shadow Institute of Technology",
                        "state": "Uttar Pradesh",
                        "parent_trust": None,
                    }
                ]
            )
        if "from public.programmes" in sql and "where institution_id" in sql:
            return FakeCursor([])
        if "join public.institutions i" in sql:
            return FakeCursor(ENROLMENT_ROWS)
        if "has_apaar" in sql:
            return FakeCursor(SHADOW_ADMISSIONS)
        if "from public.seat_pools" in sql:
            return FakeCursor(SHADOW_INTAKE)
        if "from public.integrity_scores" in sql:
            if self.score is None:
                return FakeCursor([])
            return FakeCursor(
                [{"institution_id": uuid.UUID(params[0]), "score": self.score, "tier": "high"}]
            )
        if "from public.review_cases" in sql:
            return FakeCursor([{"n": 4}])
        raise AssertionError(f"unexpected query: {sql}")


@pytest.fixture
def real_db(settings_override, monkeypatch):
    settings_override(real_domains="reports")
    conn = FakeConn()

    @contextmanager
    def fake_transaction():
        yield conn

    monkeypatch.setattr(router, "transaction", fake_transaction)
    return conn


def schema(name: str) -> dict:
    return json.loads(OPENAPI.read_text(encoding="utf-8"))["components"]["schemas"][name]


def assert_shape(body: dict, name: str) -> None:
    spec = schema(name)
    assert set(spec["required"]) <= set(body)
    assert set(body) <= set(spec["properties"])


# --------------------------------------------------------------------- fixture mode
@pytest.mark.parametrize("role", ["aicte_admin", "auditor"])
@pytest.mark.parametrize("group_by", GROUPS)
def test_aicte_roles_get_enrolment_for_every_group(client, role, group_by):
    r = client.get("/reports/enrolment", params={"group_by": group_by}, headers=as_role(role))
    assert r.status_code == 200, r.text
    assert r.headers["X-Mock"] == "true"
    body = r.json()
    assert body["group_by"] == group_by
    assert body["total"] == sum(row["admissions"] for row in body["rows"])
    assert_shape(body, "EnrolmentReport")
    for row in body["rows"]:
        assert_shape(row, "EnrolmentRow")


@pytest.mark.parametrize("role", ["institution_admin", "institution_operator"])
def test_institution_users_cannot_read_aicte_wide_enrolment(client, role):
    r = client.get("/reports/enrolment", headers=as_role(role, SHADOW))
    assert r.status_code == 403


def test_institution_admin_reads_own_accreditation_only(client):
    headers = as_role("institution_admin", SHADOW)
    own = client.get(f"/reports/accreditation/{SHADOW}", headers=headers)
    assert own.status_code == 200
    assert own.headers["X-Mock"] == "true"
    assert own.json()["institution_id"] == SHADOW
    assert client.get(f"/reports/accreditation/{NARMADA}", headers=headers).status_code == 403


def test_fixture_accreditation_matches_contract(client):
    body = client.get(f"/reports/accreditation/{SHADOW}", headers=as_role("auditor")).json()
    assert_shape(body, "AccreditationReport")
    cse = next(p for p in body["programmes"] if p["programme_code"] == "UG-CSE")
    assert (cse["admitted_total"], cse["approved_total"], cse["overrun"]) == (71, 60, 11)


@pytest.mark.parametrize(
    "path",
    ["/reports/enrolment?group_by=state", f"/reports/accreditation/{SHADOW}"],
)
def test_pdf_and_xlsx_files(client, path):
    pdf = client.get(path, params={"format": "pdf"}, headers=as_role("aicte_admin"))
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")

    xlsx = client.get(path, params={"format": "xlsx"}, headers=as_role("aicte_admin"))
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"] == XLSX
    sheet = openpyxl.load_workbook(io.BytesIO(xlsx.content)).active
    assert sheet.max_row > 1


# ------------------------------------------------------------------------ real mode
@pytest.mark.parametrize("group_by", GROUPS)
def test_real_enrolment_counts_active_admissions(client, real_db, group_by):
    r = client.get("/reports/enrolment", params={"group_by": group_by}, headers=as_role("auditor"))
    assert r.status_code == 200, r.text
    assert r.headers.get("X-Mock") != "true"
    body = r.json()
    assert_shape(body, "EnrolmentReport")
    assert body["total"] == len(ENROLMENT_ROWS)
    assert sum(row["opted_out"] for row in body["rows"]) == 1
    sql, params = real_db.calls[0]
    assert "a.active" in sql
    assert params == ("2026-27",)


def test_real_enrolment_forbidden_for_institution_users(client, real_db):
    r = client.get("/reports/enrolment", headers=as_role("institution_admin", SHADOW))
    assert r.status_code == 403
    assert real_db.calls == []


def test_real_accreditation_is_scoped_to_the_institution(client, real_db):
    r = client.get(f"/reports/accreditation/{SHADOW}", headers=as_role("institution_admin", SHADOW))
    assert r.status_code == 200, r.text
    assert r.headers.get("X-Mock") != "true"
    body = r.json()
    assert_shape(body, "AccreditationReport")
    assert body["integrity_score"] == 38.0
    assert body["integrity_note"] is None
    assert body["open_review_cases"] == 4
    assert body["apaar_coverage_pct"] == round(100 * 69 / 71, 1)
    programmes = {p["programme_code"]: p for p in body["programmes"]}
    assert programmes["UG-CSE"] == {
        "programme_code": "UG-CSE",
        "approved_total": 60,
        "admitted_total": 71,
        "overrun": 11,
    }
    assert programmes["UG-CE"]["admitted_total"] == 0
    assert programmes["UG-CE"]["overrun"] == 0
    for sql, params in real_db.calls:
        assert SHADOW in params, sql
    assert all("a.active" in sql for sql, _ in real_db.calls if "public.admissions" in sql)


def test_real_accreditation_other_institution_is_403(client, real_db):
    r = client.get(
        f"/reports/accreditation/{NARMADA}", headers=as_role("institution_admin", SHADOW)
    )
    assert r.status_code == 403
    assert real_db.calls == []


def test_real_accreditation_malformed_id_is_404(client, real_db):
    r = client.get("/reports/accreditation/not-a-uuid", headers=as_role("aicte_admin"))
    assert r.status_code == 404


def test_missing_integrity_score_still_returns_report(client, real_db):
    real_db.score = None
    r = client.get(f"/reports/accreditation/{SHADOW}", headers=as_role("aicte_admin"))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["integrity_score"] is None
    assert "not yet computed" in body["integrity_note"]
    assert body["programmes"]

    pdf = client.get(
        f"/reports/accreditation/{SHADOW}", params={"format": "pdf"}, headers=as_role("aicte_admin")
    )
    assert pdf.status_code == 200
    assert b"not yet computed" in pdf.content


@pytest.mark.parametrize(
    "path",
    ["/reports/enrolment?group_by=programme", f"/reports/accreditation/{SHADOW}"],
)
def test_real_pdf_and_xlsx_files(client, real_db, path):
    pdf = client.get(path, params={"format": "pdf"}, headers=as_role("aicte_admin"))
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content.startswith(b"%PDF")
    assert pdf.headers.get("X-Mock") != "true"

    xlsx = client.get(path, params={"format": "xlsx"}, headers=as_role("aicte_admin"))
    assert xlsx.status_code == 200
    assert xlsx.headers["content-type"] == XLSX
    assert openpyxl.load_workbook(io.BytesIO(xlsx.content)).active.max_row > 1
